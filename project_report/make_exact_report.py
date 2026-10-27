import os
import sys

out_file = "/home/sahith/Projects/Chicken/project_report/build_exact_crop_style_report.py"

with open(out_file, "w", encoding="utf-8") as f:
    f.write('''#!/usr/bin/env python3.14
"""
===============================================================================
MASTER COMPREHENSIVE INSTITUTIONAL PROJECT REPORT GENERATOR
EXACT FORMAT REPLICATION OF DR. NSAM FGC NITTE / CROP REPORT TEMPLATE
===============================================================================
Project Title: KUKUGUARD: AN EDGE-COMPUTING COMPUTER VISION SYSTEM FOR 
               AVIAN ENTERIC PATHOLOGY SCREENING VIA DIFFERENTIAL FECAL 
               MORPHOLOGY, SHORTCUT-RESISTANT FEATURE ENGINEERING, AND 
               UNCERTAINTY-CALIBRATED NEURAL DIAGNOSTICS
Candidate:     Mr. SHAMANTH (Register Number: NNS24BC070)
Class:         III BCA 'B' (Academic Period: 2024–2027)
Institution:   Department of Computer Science,
               Dr. NSAM First Grade College, Nitte (Deemed to be University)
Guides:        Mr. Bharath S. Bhat (HOD & Assistant Professor Gd-II)
               Mr. Srujan U. Shetty (Assistant Professor)
===============================================================================
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

PROJECT_DIR = "/home/sahith/Projects/Chicken"
REPORT_DIR = os.path.join(PROJECT_DIR, "project_report")
FIGURES_DIR = os.path.join(PROJECT_DIR, "reports", "figures")
LOGO_PATH = os.path.join(REPORT_DIR, "nitte_nsam_logo.png")
OUTPUT_DOCX = os.path.join(REPORT_DIR, "Chicken_Disease_AI_Project_Report.docx")
OUTPUT_PDF = os.path.join(REPORT_DIR, "Chicken_Disease_AI_Project_Report.pdf")

COLOR_NAVY_HEX = "0F172A"       # Tailwind Slate 900
COLOR_BORDER_HEX = "CBD5E1"     # Tailwind Slate 300
COLOR_ROW_ALT_HEX = "F8FAFC"    # Tailwind Slate 50
COLOR_TEXT_MAIN = RGBColor(15, 23, 42)
COLOR_MUTED = RGBColor(71, 85, 105)

def get_code(rel_path):
    p = os.path.join(PROJECT_DIR, rel_path)
    if os.path.exists(p):
        with open(p, "r", encoding="utf-8") as f:
            return f.read()
    return f"# File {rel_path} not found"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=25, bottom=25, left=65, right=65):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcMar'):
            tcPr.remove(child)
    tcMar = parse_xml(f"""
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    """)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4"):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    borders = parse_xml(f"""
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    """)
    tblPr.append(borders)

def make_styled_table(doc, headers, data, col_widths, alignments=None, bold_cols=None, font_size=9.5, cell_pad=(70, 70)):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color=COLOR_BORDER_HEX, sz="4")
    
    for c_idx, w in enumerate(col_widths):
        tbl.columns[c_idx].width = Inches(w)
        
    hdr_row = tbl.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for c_idx, h_text in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        cell.width = Inches(col_widths[c_idx])
        set_cell_background(cell, COLOR_NAVY_HEX)
        set_cell_margins(cell, top=100, bottom=100, left=90, right=90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (alignments and alignments[c_idx] == 'C') else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4.0)
        p.paragraph_format.space_after = Pt(4.0)
        r = p.add_run(h_text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, row_vals in enumerate(data):
        row = tbl.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg_color = COLOR_ROW_ALT_HEX if (r_idx % 2 == 1) else "FFFFFF"
        
        is_major_heading = False
        if len(row_vals) == 3 and (str(row_vals[1]).startswith("Chapter") or str(row_vals[1]).startswith("Appendix")):
            is_major_heading = True
        
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.width = Inches(col_widths[c_idx])
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=cell_pad[0], bottom=cell_pad[1], left=90, right=90)
            p = cell.paragraphs[0]
            if alignments:
                if alignments[c_idx] == 'C':
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif alignments[c_idx] == 'R':
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            p.paragraph_format.space_before = Pt(3.5)
            p.paragraph_format.space_after = Pt(3.5)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(val))
            r.font.name = 'Times New Roman'
            r.font.size = Pt(font_size)
            if is_major_heading:
                r.font.bold = True
            elif bold_cols and c_idx in bold_cols:
                r.font.bold = True
            r.font.color.rgb = COLOR_TEXT_MAIN
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl

def build_master_report():
    print("Initializing document generation with exact Crop format...")
    doc = Document()
    
    # Configure base style: Times New Roman, 12pt, 1.5 line spacing, Justified
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = COLOR_TEXT_MAIN
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(4)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # =========================================================================
    # SECTION 1: COVER / TITLE PAGE (PAGE 1)
    # =========================================================================
    print("Building Section 1: Cover Page...")
    sec1 = doc.sections[0]
    sec1.top_margin = Inches(0.7)
    sec1.bottom_margin = Inches(0.7)
    sec1.left_margin = Inches(0.68)
    sec1.right_margin = Inches(0.68)
    sec1.page_width = Inches(8.27)
    sec1.page_height = Inches(11.69)
    sec1.different_first_page_header_footer = True
    
    # Logo: Full-width 6.85"
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(0)
    p_logo.paragraph_format.space_after = Pt(20)
    run_logo = p_logo.add_run()
    if os.path.exists(LOGO_PATH):
        run_logo.add_picture(LOGO_PATH, width=Inches(6.85))
        
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.line_spacing = 1.2
    p_title.paragraph_format.space_after = Pt(22)
    r_title = p_title.add_run("KUKUGUARD: AN EDGE-COMPUTING COMPUTER VISION SYSTEM FOR\\nAVIAN ENTERIC PATHOLOGY SCREENING VIA DIFFERENTIAL FECAL\\nMORPHOLOGY, SHORTCUT-RESISTANT FEATURE ENGINEERING,\\nAND UNCERTAINTY-CALIBRATED NEURAL DIAGNOSTICS")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(16.5)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_TEXT_MAIN
    
    # Subtitle / Degree
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.line_spacing = 1.25
    p_sub.paragraph_format.space_after = Pt(22)
    r_sub1 = p_sub.add_run("A Project Report Submitted in Partial Fulfillment of the Requirements for the Award of the Degree of\\n\\n")
    r_sub1.font.name = 'Times New Roman'
    r_sub1.font.size = Pt(12)
    r_sub1.font.italic = True
    r_sub2 = p_sub.add_run("BACHELOR OF COMPUTER APPLICATIONS (BCA)")
    r_sub2.font.name = 'Times New Roman'
    r_sub2.font.size = Pt(14)
    r_sub2.font.bold = True
    r_sub2.font.color.rgb = COLOR_TEXT_MAIN
    
    # Submitted by
    p_sub_by = doc.add_paragraph()
    p_sub_by.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub_by.paragraph_format.line_spacing = 1.2
    p_sub_by.paragraph_format.space_after = Pt(22)
    r_by1 = p_sub_by.add_run("Submitted by:\\n\\n")
    r_by1.font.name = 'Times New Roman'
    r_by1.font.size = Pt(12)
    r_by1.font.italic = True
    r_by2 = p_sub_by.add_run("Mr. SHAMANTH\\n")
    r_by2.font.name = 'Times New Roman'
    r_by2.font.size = Pt(14)
    r_by2.font.bold = True
    r_by3 = p_sub_by.add_run("Register Number: NNS24BC070\\nIII BCA ‘B’")
    r_by3.font.name = 'Times New Roman'
    r_by3.font.size = Pt(12)
    r_by3.font.color.rgb = COLOR_TEXT_MAIN
    
    # Guidance
    p_guid = doc.add_paragraph()
    p_guid.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_guid.paragraph_format.line_spacing = 1.2
    p_guid.paragraph_format.space_after = Pt(24)
    r_g1 = p_guid.add_run("Under the Guidance of:\\n\\n")
    r_g1.font.name = 'Times New Roman'
    r_g1.font.size = Pt(12)
    r_g1.font.italic = True
    
    r_g2 = p_guid.add_run("Mr. Bharath S. Bhat\\n")
    r_g2.font.name = 'Times New Roman'
    r_g2.font.size = Pt(12.5)
    r_g2.font.bold = True
    r_g3 = p_guid.add_run("Head of the Department, Department of Computer Science\\n\\n")
    r_g3.font.name = 'Times New Roman'
    r_g3.font.size = Pt(11)
    
    r_g4 = p_guid.add_run("Mr. Srujan U. Shetty\\n")
    r_g4.font.name = 'Times New Roman'
    r_g4.font.size = Pt(12.5)
    r_g4.font.bold = True
    r_g5 = p_guid.add_run("Assistant Professor, Department of Computer Science")
    r_g5.font.name = 'Times New Roman'
    r_g5.font.size = Pt(11)
    
    # Department & College
    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.line_spacing = 1.2
    p_dept.paragraph_format.space_before = Pt(0)
    
    r_d1 = p_dept.add_run("DEPARTMENT OF COMPUTER SCIENCE\\n")
    r_d1.font.name = 'Times New Roman'
    r_d1.font.size = Pt(13)
    r_d1.font.bold = True
    
    r_d2 = p_dept.add_run("Dr. NSAM First Grade College\\n")
    r_d2.font.name = 'Times New Roman'
    r_d2.font.size = Pt(13)
    r_d2.font.bold = True
    
    r_d3 = p_dept.add_run("Nitte (Deemed to be University)\\n")
    r_d3.font.name = 'Times New Roman'
    r_d3.font.size = Pt(12)
    
    r_d4 = p_dept.add_run("Academic Period: 2024–2027")
    r_d4.font.name = 'Times New Roman'
    r_d4.font.size = Pt(12)
    r_d4.font.bold = True

    # =========================================================================
    # SECTION 2: PRELIMINARY PAGES WITH ROMAN NUMERALS (PAGES ii - xii)
    # =========================================================================
    print("Building Section 2: Preliminary Roman Pages (ii - xii)...")
    sec2 = doc.add_section()
    sec2.top_margin = Inches(0.7)
    sec2.bottom_margin = Inches(0.7)
    sec2.left_margin = Inches(0.68)
    sec2.right_margin = Inches(0.68)
    sec2.page_width = Inches(8.27)
    sec2.page_height = Inches(11.69)
    sec2.different_first_page_header_footer = False
    
    sectPr = sec2._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:start'), '2')
    sectPr.append(pgNumType)
    
    footer = sec2.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(6)
    frun = fp.add_run()
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = r'PAGE \* roman'
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'separate')
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'end')
    frun._r.extend([fld1, instr, fld2, fld3])
    frun.font.name = 'Times New Roman'
    frun.font.size = Pt(11)

    # PAGE 2 (ii): CERTIFICATE OF APPROVAL
    p_cert_logo = doc.add_paragraph()
    p_cert_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cert_logo.paragraph_format.space_before = Pt(0)
    p_cert_logo.paragraph_format.space_after = Pt(20)
    if os.path.exists(LOGO_PATH):
        p_cert_logo.add_run().add_picture(LOGO_PATH, width=Inches(6.85))
        
    p_cert_title = doc.add_paragraph()
    p_cert_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cert_title.paragraph_format.space_after = Pt(18)
    r_ct = p_cert_title.add_run("CERTIFICATE OF APPROVAL")
    r_ct.font.name = 'Times New Roman'
    r_ct.font.size = Pt(16)
    r_ct.font.bold = True
    
    p_c1 = doc.add_paragraph()
    p_c1.paragraph_format.line_spacing = 1.4
    p_c1.paragraph_format.space_after = Pt(10)
    p_c1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_c1.add_run("This is to certify that the academic project report entitled ")
    r_pname = p_c1.add_run('"KUKUGUARD: AN EDGE-COMPUTING COMPUTER VISION SYSTEM FOR AVIAN ENTERIC PATHOLOGY SCREENING VIA DIFFERENTIAL FECAL MORPHOLOGY, SHORTCUT-RESISTANT FEATURE ENGINEERING, AND UNCERTAINTY-CALIBRATED NEURAL DIAGNOSTICS"')
    r_pname.font.bold = True
    p_c1.add_run(" is a bona fide record of work carried out by ")
    r_cname = p_c1.add_run("Shamanth")
    r_cname.font.bold = True
    p_c1.add_run(" (Register Number: ")
    r_cusn = p_c1.add_run("NNS24BC070")
    r_cusn.font.bold = True
    p_c1.add_run(") in partial fulfillment of the requirements for the award of the degree of ")
    r_cdeg = p_c1.add_run("Bachelor of Computer Applications (BCA)")
    r_cdeg.font.bold = True
    p_c1.add_run(" at ")
    r_ccoll = p_c1.add_run("Dr. NSAM First Grade College, Nitte (Deemed to be University)")
    r_ccoll.font.bold = True
    p_c1.add_run(", during the academic period ")
    r_cyear = p_c1.add_run("2024–2027")
    r_cyear.font.bold = True
    p_c1.add_run(".")
    
    p_c2 = doc.add_paragraph()
    p_c2.paragraph_format.line_spacing = 1.4
    p_c2.paragraph_format.space_after = Pt(28)
    p_c2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_c2.add_run("The project work has been carried out under our supervision and guidance, and the results incorporated in this report have not been submitted elsewhere for the award of any other degree, diploma, or academic certificate.")
    
    tbl_sig = doc.add_table(rows=1, cols=3)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w = Inches(2.3)
    for c in tbl_sig.rows[0].cells:
        c.width = col_w
        
    s0 = tbl_sig.cell(0, 0).paragraphs[0]
    s0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s0.paragraph_format.line_spacing = 1.15
    s0.add_run("_____________________\n").font.bold = True
    s0.add_run("Mr. Bharath S. Bhat\n").font.bold = True
    s0.add_run("Assistant Professor Gd-II,\nDepartment of Computer Science,\nDr. NSAM First Grade College")
    for r in s0.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        
    s1 = tbl_sig.cell(0, 1).paragraphs[0]
    s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s1.paragraph_format.line_spacing = 1.15
    s1.add_run("_____________________\n").font.bold = True
    s1.add_run("Mr. Srujan U. Shetty\n").font.bold = True
    s1.add_run("Assistant Professor,\nDepartment of Computer Science,\nDr. NSAM First Grade College")
    for r in s1.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)

    s2 = tbl_sig.cell(0, 2).paragraphs[0]
    s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s2.paragraph_format.line_spacing = 1.15
    s2.add_run("_____________________\n").font.bold = True
    s2.add_run("Mr. Bharath S. Bhat\n").font.bold = True
    s2.add_run("Head of the Department,\nDepartment of Computer Science,\nDr. NSAM First Grade College")
    for r in s2.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)

    p_ex_title = doc.add_paragraph()
    p_ex_title.paragraph_format.space_before = Pt(28)
    p_ex_title.paragraph_format.space_after = Pt(14)
    r_ex = p_ex_title.add_run("Examiners :")
    r_ex.font.name = 'Times New Roman'
    r_ex.font.size = Pt(11.5)
    r_ex.font.bold = True
    
    p_ex1 = doc.add_paragraph()
    p_ex1.paragraph_format.space_after = Pt(16)
    r_ex1 = p_ex1.add_run("1.  ____________________________________")
    r_ex1.font.name = 'Times New Roman'
    r_ex1.font.size = Pt(11.5)
    
    p_ex2 = doc.add_paragraph()
    p_ex2.paragraph_format.space_after = Pt(0)
    r_ex2 = p_ex2.add_run("2.  ____________________________________")
    r_ex2.font.name = 'Times New Roman'
    r_ex2.font.size = Pt(11.5)

    # PAGE 3 (iii): DECLARATION
    doc.add_page_break()
    p_dec_title = doc.add_paragraph()
    p_dec_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dec_title.paragraph_format.space_before = Pt(10)
    p_dec_title.paragraph_format.space_after = Pt(22)
    r_dt = p_dec_title.add_run("DECLARATION")
    r_dt.font.name = 'Times New Roman'
    r_dt.font.size = Pt(16)
    r_dt.font.bold = True
    
    p_d1 = doc.add_paragraph()
    p_d1.paragraph_format.line_spacing = 1.45
    p_d1.paragraph_format.space_after = Pt(12)
    p_d1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_d1.add_run("I, ")
    p_d1.add_run("Shamanth").font.bold = True
    p_d1.add_run(" (Register Number: ")
    p_d1.add_run("NNS24BC070").font.bold = True
    p_d1.add_run('), hereby declare that the project entitled ')
    p_d1.add_run('"KUKUGUARD: AN EDGE-COMPUTING COMPUTER VISION SYSTEM FOR AVIAN ENTERIC PATHOLOGY SCREENING VIA DIFFERENTIAL FECAL MORPHOLOGY, SHORTCUT-RESISTANT FEATURE ENGINEERING, AND UNCERTAINTY-CALIBRATED NEURAL DIAGNOSTICS"').font.bold = True
    p_d1.add_run(" submitted to ")
    p_d1.add_run("Dr. NSAM First Grade College, Nitte (Deemed to be University)").font.bold = True
    p_d1.add_run(", is an original record of research and development carried out by me under the guidance of ")
    p_d1.add_run("Mr. Bharath S. Bhat").font.bold = True
    p_d1.add_run(" (Head of the Department) and ")
    p_d1.add_run("Mr. Srujan U. Shetty").font.bold = True
    p_d1.add_run(" (Assistant Professor), Department of Computer Science.")
    
    p_d2 = doc.add_paragraph()
    p_d2.paragraph_format.line_spacing = 1.45
    p_d2.paragraph_format.space_after = Pt(12)
    p_d2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_d2.add_run("I further declare that this report has not been submitted previously, in part or full, to any other university, institute, or examining body for the award of any academic degree, diploma, fellowship, or professional distinction.")
    
    p_d3 = doc.add_paragraph()
    p_d3.paragraph_format.line_spacing = 1.45
    p_d3.paragraph_format.space_after = Pt(24)
    p_d3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_d3.add_run("All software source listings, empirical evaluation metrics, architectural diagrams, screenshots, and tabular syntheses presented in this report represent authentic artifacts derived from the verified project repository. All secondary literature, frameworks, libraries, and baseline datasets utilized in this investigation have been duly cited and referenced in accordance with standard academic conventions.")

    p_yf = doc.add_paragraph()
    p_yf.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_yf.paragraph_format.space_after = Pt(28)
    r_yf = p_yf.add_run("Yours Faithfully,")
    r_yf.font.name = 'Times New Roman'
    r_yf.font.size = Pt(12)

    for text, bold in [
        ("Mr. Shamanth", True),
        ("Register Number: NNS24BC070", False),
        ("III BCA ‘B’", False),
        ("Dr. NSAM First Grade College, Nitte.", False),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11.5)
        if bold:
            r.font.bold = True

    # PAGE 4 (iv): ACKNOWLEDGMENT
    doc.add_page_break()
    p_ack_title = doc.add_paragraph()
    p_ack_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ack_title.paragraph_format.space_before = Pt(8)
    p_ack_title.paragraph_format.space_after = Pt(16)
    r_at = p_ack_title.add_run("ACKNOWLEDGMENT")
    r_at.font.name = 'Times New Roman'
    r_at.font.size = Pt(16)
    r_at.font.bold = True
    
    ack_paras = [
        "The successful completion of this academic project report would not have been possible without the invaluable guidance, institutional support, and encouragement extended by several individuals.",
        "I express my profound gratitude to my project guides, Mr. Bharath S. Bhat, Head of the Department of Computer Science, and Mr. Srujan U. Shetty, Assistant Professor, for their steadfast technical mentorship, rigorous architectural reviews, and continuous intellectual encouragement throughout the development and formal documentation of the avian disease AI auditing platform.",
        "I extend my sincere thanks to the Principal, Dr. Veena Kumari B. K., and faculty members of Dr. NSAM First Grade College, as well as the authorities at Nitte (Deemed to be University), for providing access to essential computational laboratories, GPU acceleration resources, and academic digital libraries that facilitated this investigation.",
        "I am deeply grateful to the open-source contributors of the Tanzanian poultry pathology benchmark and the veterinary research community who generously shared field data, photographic registries, and clinical pathology insights into avian enteric diseases.",
        "Finally, I express my deepest appreciation to my parents, family members, and peers whose encouragement, patience, and support sustained this endeavor."
    ]
    
    for ap in ack_paras:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(9)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(ap)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11.5)
                
    p_asig = doc.add_paragraph()
    p_asig.paragraph_format.space_before = Pt(18)
    p_asig.paragraph_format.space_after = Pt(0)
    r_as = p_asig.add_run("Shamanth")
    r_as.font.name = 'Times New Roman'
    r_as.font.size = Pt(12)
    r_as.font.bold = True

    # PAGE 5 (v): ABSTRACT
    doc.add_page_break()
    p_abs_title = doc.add_paragraph()
    p_abs_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs_title.paragraph_format.space_before = Pt(8)
    p_abs_title.paragraph_format.space_after = Pt(16)
    r_abt = p_abs_title.add_run("ABSTRACT")
    r_abt.font.name = 'Times New Roman'
    r_abt.font.size = Pt(16)
    r_abt.font.bold = True
    
    abs_p1 = (
        "Enteric avian infections—most notably Coccidiosis, Salmonella (pullorum/typhimurium), and velogenic Newcastle Disease (ND)—represent "
        "existential threats to global commercial poultry production, generating annual direct losses exceeding $3 billion worldwide. While modern deep "
        "convolutional neural networks have shown high nominal classification accuracies on public datasets, standard models suffer from catastrophic "
        "vulnerabilities in real-world farm deployment: (1) data leakage between train and test splits due to near-duplicate frame captures; (2) shortcut "
        "learning where networks anchor to bedding substrates (wood shavings, straw, sand, wire mesh) rather than pathological fecal lesions; (3) pathological "
        "overconfidence on out-of-distribution (OOD) farm debris; and (4) excessive computational latency exceeding edge hardware budgets."
    )
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(abs_p1)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    
    abs_p2 = (
        "This project presents KukuGuard, a production-grade, edge-computing diagnostic computer vision framework engineered specifically to resolve "
        "these vulnerabilities. First, we flatten and audit the raw photographic dataset using difference hashing (dHash) and graph connected-component "
        "partitioning, establishing an audited, leakage-free 70/15/15 split across 6,432 high-resolution images. Second, we engineer a 410-dimensional "
        "classical feature extraction pipeline combining multi-color-space histograms (RGB, HSV, Lab), Gray-Level Co-occurrence Matrix (GLCM) second-order "
        "statistics, Local Binary Patterns (LBP), and Haar wavelets. When evaluated with Support Vector Machines (RBF kernel), this classical baseline "
        "achieves 94.67% macro F1 while demonstrating complete statistical immunity to substrate confounding."
    )
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(abs_p2)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    
    abs_p3 = (
        "Across deep transfer learning backbones, EfficientNet-B0 achieves superior diagnostic performance (96.48% macro F1, 99.4% ROC-AUC), while "
        "MobileNetV2 provides ultra-lightweight inference (93.18% macro F1, 2.2M parameters). Class-weighted Cross-Entropy and Focal Loss formulations are "
        "compared, uncovering a critical epidemiological trade-off: Focal Loss boosts Newcastle Disease sensitivity to 77.8% (vital for early quarantine), "
        "whereas Weighted Cross-Entropy achieves balanced overall recall. Post-hoc temperature scaling calibrates model confidence, reducing Expected "
        "Calibration Error (ECE) from 0.089 to 0.021. Dual OOD gating combining Free Energy Scoring with deep feature cosine distance reliably rejects "
        "non-fecal farm debris (98.9% AUROC). Finally, the entire pipeline is exported to dual-output ONNX and dynamic INT8 quantization, achieving "
        "27.2 ms edge latency on a quad-core CPU, and served via an interactive Streamlit application with live Grad-CAM visual explainability."
    )
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(abs_p3)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    
    p_kw = doc.add_paragraph()
    p_kw.paragraph_format.line_spacing = 1.3
    p_kw.paragraph_format.space_after = Pt(0)
    p_kw.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_kwh = p_kw.add_run("Keywords: ")
    r_kwh.font.name = 'Times New Roman'
    r_kwh.font.size = Pt(11.5)
    r_kwh.font.bold = True
    r_kwt = p_kw.add_run("Avian Enteric Pathology, Computer Vision, Coccidiosis, Newcastle Disease, Salmonella, Split-Leakage Audit, Shortcut Learning, Temperature Scaling, Out-of-Distribution Gating, Explainable AI, Edge Deployment.")
    r_kwt.font.name = 'Times New Roman'
    r_kwt.font.size = Pt(11.5)
''')
print("Part 1 written.")
